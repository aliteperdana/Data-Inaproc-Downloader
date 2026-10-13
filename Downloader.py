import os
import re
import sys
import glob
import time
import logging
from datetime import datetime
from urllib.parse import urlparse, parse_qs

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager

# ---------------------------------------------------------------
# Lokasi folder: semua relatif terhadap lokasi file skrip ini,
# jadi hasilnya sama walau dijalankan dari Task Scheduler/terminal lain
# ---------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DOWNLOAD_DIR = os.path.join(BASE_DIR, 'Downloads')
LOG_DIR = os.path.join(BASE_DIR, 'Logs')
LINKS_FILE = os.path.join(BASE_DIR, 'links.txt')

TIMEOUT_HALAMAN = 30     # detik, menunggu elemen muncul
TIMEOUT_UNDUHAN = 120    # detik, menunggu file selesai diunduh

KONVERSI_XLSX = True     # True: buat file .xlsx dari setiap CSV yang diunduh
SIMPAN_CSV = True        # False: hapus CSV setelah berhasil dikonversi ke .xlsx

BULAN = {
    'jan': '01', 'feb': '02', 'mar': '03', 'apr': '04', 'mei': '05', 'jun': '06',
    'jul': '07', 'agu': '08', 'ags': '08', 'sep': '09', 'okt': '10', 'nov': '11', 'des': '12',
}

os.makedirs(DOWNLOAD_DIR, exist_ok=True)
os.makedirs(LOG_DIR, exist_ok=True)

# ---------------------------------------------------------------
# Logging: file baru setiap skrip dijalankan + tampil di layar
# ---------------------------------------------------------------
log_path = os.path.join(LOG_DIR, f"log_{datetime.now():%Y%m%d_%H%M%S}.txt")
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s [%(levelname)s] %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S',
    handlers=[
        logging.FileHandler(log_path, encoding='utf-8'),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger('downloader')


def bersihkan_nama(teks):
    """Ubah teks bebas menjadi aman untuk nama file (spasi dan simbol jadi _)."""
    return re.sub(r'[^A-Za-z0-9_-]+', '_', teks).strip('_')


def baca_daftar_link(path):
    """
    Satu link per baris. Baris kosong dan baris berawalan # diabaikan.
    Format yang didukung:
        https://...                  -> nama file otomatis dari link
        Nama Sendiri | https://...   -> nama file memakai 'Nama Sendiri'
    Mengembalikan daftar (label, url); label None jika tidak ditulis.
    """
    if not os.path.exists(path):
        log.error("File daftar link tidak ditemukan: %s", path)
        return []
    daftar = []
    with open(path, encoding='utf-8') as f:
        for baris in f:
            baris = baris.strip()
            if not baris or baris.startswith('#'):
                continue
            if '|' in baris:
                label, url = (x.strip() for x in baris.split('|', 1))
                daftar.append((bersihkan_nama(label) or None, url))
            else:
                daftar.append((None, baris))
    return daftar


def nama_dasar_dari_url(url):
    """
    Nama otomatis dari link, memakai parameter yang membedakan halaman.
    Contoh: .../afirmasi-pdn?tahun=2026&instansi=D376&view_table=Realisasi
            -> afirmasi-pdn_2026_D376_Realisasi
    """
    u = urlparse(url)
    bagian = [u.path.strip('/').split('/')[-1] or 'data']
    q = parse_qs(u.query)
    for kunci in ('tahun', 'instansi', 'view_table'):
        if kunci in q:
            bagian.append(q[kunci][0])
    return bersihkan_nama('_'.join(bagian))


def waktu_sinkron_ke_nama(teks):
    """'Data Tersinkron: 13 Okt 26, 02.27 WIB' -> '20261013_0227'."""
    m = re.search(r'(\d{1,2})\s+([A-Za-z]+)\s+(\d{2,4}),?\s+(\d{1,2})[.:](\d{2})', teks)
    if m:
        hari, bulan, tahun, jam, menit = m.groups()
        bln = BULAN.get(bulan[:3].lower())
        if bln:
            if len(tahun) == 2:
                tahun = '20' + tahun
            return f"{tahun}{bln}{int(hari):02d}_{int(jam):02d}{menit}"
    return re.sub(r'[^A-Za-z0-9]+', '_', teks).strip('_')


def nama_unik(direktori, nama_dasar, ekstensi):
    """
    Kembalikan path yang belum dipakai. Jika sudah ada file dengan nama sama
    (misalnya data belum disinkron ulang), tambahkan _2, _3, dst. agar tidak menimpa.
    Ekstensi .csv dan .xlsx dicek bersamaan supaya pasangannya selalu bernomor sama.
    """
    n = 1
    while True:
        akhiran = '' if n == 1 else f'_{n}'
        csv_path = os.path.join(direktori, f"{nama_dasar}{akhiran}.csv")
        xlsx_path = os.path.join(direktori, f"{nama_dasar}{akhiran}.xlsx")
        if not os.path.exists(csv_path) and not os.path.exists(xlsx_path):
            return os.path.join(direktori, f"{nama_dasar}{akhiran}{ekstensi}")
        n += 1


def baca_csv(path):
    """Baca CSV dengan deteksi pemisah otomatis dan fallback encoding."""
    for enc in ('utf-8-sig', 'latin-1'):
        try:
            df = pd.read_csv(path, sep=None, engine='python', encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Encoding CSV tidak dikenali")
    # Kolom kode (mis. Kode RUP) dibaca ulang sebagai teks agar tidak berubah jadi angka/ilmiah
    kolom_kode = [c for c in df.columns if 'kode' in str(c).lower()]
    if kolom_kode:
        df = pd.read_csv(path, sep=None, engine='python', encoding=enc,
                         dtype={c: str for c in kolom_kode})
    return df


def konversi_ke_xlsx(csv_path):
    """Konversi CSV ke XLSX (header tebal, baris header dibekukan, lebar kolom menyesuaikan)."""
    xlsx_path = os.path.splitext(csv_path)[0] + '.xlsx'
    df = baca_csv(csv_path)
    with pd.ExcelWriter(xlsx_path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Data')
        ws = writer.sheets['Data']
        for sel in ws[1]:
            sel.font = Font(bold=True)
        ws.freeze_panes = 'A2'
        for i, kolom in enumerate(df.columns, start=1):
            lebar = max([len(str(kolom))] + [len(str(v)) for v in df[kolom].head(200)])
            ws.column_dimensions[get_column_letter(i)].width = min(lebar + 2, 60)
    return xlsx_path, len(df)


def tunggu_unduhan(direktori, timeout=TIMEOUT_UNDUHAN):
    """Menunggu CSV baru muncul. Mengembalikan path file baru, atau None jika timeout."""
    awal = set(glob.glob(os.path.join(direktori, '*.csv')))
    for _ in range(timeout):
        baru = set(glob.glob(os.path.join(direktori, '*.csv'))) - awal
        sedang = glob.glob(os.path.join(direktori, '*.crdownload'))
        if baru and not sedang:
            time.sleep(1)
            return max(baru, key=os.path.getmtime)
        time.sleep(1)
    return None


def buat_driver():
    options = webdriver.ChromeOptions()
    options.add_argument('--headless')
    options.add_argument('--no-sandbox')
    options.add_argument('--disable-dev-shm-usage')
    options.add_argument('--window-size=1920,1080')
    options.add_argument('--user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/114.0.0.0 Safari/537.36')
    options.add_experimental_option("prefs", {
        "download.default_directory": DOWNLOAD_DIR,
        "download.prompt_for_download": False,
        "download.directory_upgrade": True,
    })
    driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)
    driver.execute_cdp_cmd('Browser.setDownloadBehavior', {
        'behavior': 'allow',
        'downloadPath': DOWNLOAD_DIR,
    })
    return driver


def proses_satu_link(driver, url, label=None):
    """Unduh CSV dari satu link. Mengembalikan path file akhir, atau None jika gagal."""
    wait = WebDriverWait(driver, TIMEOUT_HALAMAN)
    driver.get(url)

    # Ambil waktu sinkron dari badge (tidak fatal jika gagal)
    teks_sinkron = None
    try:
        badge = wait.until(EC.presence_of_element_located(
            (By.XPATH, "//*[contains(text(), 'Data Tersinkron')]")
        ))
        teks_sinkron = badge.text.strip()
        log.info("Badge: %s", teks_sinkron)
    except Exception:
        log.warning("Badge 'Data Tersinkron' tidak ditemukan, memakai waktu saat ini untuk nama file.")

    tombol = wait.until(EC.element_to_be_clickable(
        (By.XPATH, "//button[contains(text(), 'Download CSV')]")
    ))
    tombol.click()
    log.info("Tombol Download CSV diklik, menunggu file...")

    file_baru = tunggu_unduhan(DOWNLOAD_DIR)
    if not file_baru:
        log.error("Timeout (%d detik), file tidak muncul.", TIMEOUT_UNDUHAN)
        return None

    stempel = waktu_sinkron_ke_nama(teks_sinkron) if teks_sinkron else datetime.now().strftime('%Y%m%d_%H%M')
    nama_dasar = f"{label or nama_dasar_dari_url(url)}_sinkron_{stempel}"
    tujuan = nama_unik(DOWNLOAD_DIR, nama_dasar, '.csv')  # tidak menimpa file yang sudah ada
    os.rename(file_baru, tujuan)
    log.info("Tersimpan: %s", os.path.basename(tujuan))

    if KONVERSI_XLSX:
        try:
            xlsx_path, jumlah_baris = konversi_ke_xlsx(tujuan)
            log.info("Dikonversi ke XLSX: %s (%d baris data)", os.path.basename(xlsx_path), jumlah_baris)
            if not SIMPAN_CSV:
                os.remove(tujuan)
                log.info("CSV dihapus: %s", os.path.basename(tujuan))
                return xlsx_path
        except Exception:
            log.exception("Konversi ke XLSX gagal, file CSV tetap disimpan")
    return tujuan


def main():
    log.info("=== Downloader dimulai ===")
    log.info("File log: %s", log_path)

    links = baca_daftar_link(LINKS_FILE)
    if not links:
        log.error("Tidak ada link untuk diproses. Isi file: %s", LINKS_FILE)
        return 1
    log.info("Jumlah link: %d", len(links))

    try:
        driver = buat_driver()
    except Exception:
        log.exception("Gagal membuka Chrome. Pastikan Google Chrome terpasang (lihat README, bagian Instalasi).")
        return 3

    berhasil, gagal = 0, 0
    try:
        for i, (label, url) in enumerate(links, start=1):
            log.info("--- [%d/%d] %s", i, len(links), url)
            try:
                hasil = proses_satu_link(driver, url, label)
            except Exception:
                log.exception("Gagal memproses link ini")
                hasil = None
            if hasil:
                berhasil += 1
            else:
                gagal += 1
    finally:
        driver.quit()

    log.info("=== Selesai. Berhasil: %d, gagal: %d ===", berhasil, gagal)
    return 0 if gagal == 0 else 2


if __name__ == '__main__':
    sys.exit(main())